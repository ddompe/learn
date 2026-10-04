"""Convert an xlsx and a PDF to Markdown by hand, to see where conversion gets hard."""

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

DATA = Path(__file__).resolve().parent.parent / "data"

sheet = load_workbook(DATA / "cafe_central_monthly_summary.xlsx")["Resumen"]
print("Merged headers:", sorted(str(r) for r in sheet.merged_cells.ranges))
rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=3)]
header, body = rows[0], rows[1:]
print("| " + " | ".join(map(str, header)) + " |")
print("| " + " | ".join("---" for _ in header) + " |")
for row in body:
    print("| " + " | ".join(map(str, row)) + " |")
print("(The two 'Ventas' and two 'Monto' headers lost their month: it sat in merged cells.)")

pdf = (DATA / "cafe_central_invoice.pdf").read_bytes().decode("latin-1")
lines = defaultdict(list)
for x, y, text in re.findall(r"1 0 0 1 (\d+) (\d+) Tm \(([^)]*)\) Tj", pdf):
    lines[int(y)].append((int(x), text))
print("\nInvoice rows rebuilt by grouping text with the same height:")
for y in (630, 614, 598):
    print(" | ".join(text for _, text in sorted(lines[y])))
